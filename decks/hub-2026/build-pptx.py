#!/usr/bin/env python3
"""Sklada deck.md bezposrednio w PowerPointa - z pominieciem PDF-a.

Dlaczego nie z PDF-a: rozmiar tekstu koduje w tym decku role slajdu (tytulowy 118px,
koncowy 96px, rzadki 64px, zwykly 41px, podpis rysunku 31px). W PDF ta informacja nie
istnieje, zostaja same glify, wiec konwerter mapuje kazdy rozmiar doslownie i typografia
sie rozjezdza. Tutaj czytamy te role wprost z deck.md przez deck_source.py.

Geometria przenosi sie ze slajdu 1600x900 px na slajd 16:9 dokladnie:
    12192000 EMU / 1600 px = 7620        6858000 EMU / 900 px = 7620
czyli 1 px CSS = 7620 EMU, bez reszty w obu osiach. Rozmiary fontow:
    pt = px * 0.85 (--fs) * 0.6 (72pt / 120px na cal)
"""
import re, sys, pathlib
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import deck_source as ds

HERE = pathlib.Path(__file__).parent
OUT  = HERE / "Frontier-u-siebie.pptx"

# ---------- skala ----------
PXE = 7620          # EMU na piksel CSS
FS  = 0.85          # --fs z deck-print.css
def E(px):  return Emu(int(round(px * PXE)))
def P(px):  return Pt(round(px * FS * 0.6, 2))
def SZ(px): return px * FS                      # rozmiar renderowany w px plotna

W, H = 1600, 900
PAD_T, PAD_X, PAD_B = 58, 72, 56
CONTENT_W = W - 2 * PAD_X                       # 1456

# ---------- kolory z :root ----------
TEXT   = RGBColor(0x0A, 0x0A, 0x0A)
MUTED  = RGBColor(0x3F, 0x3F, 0x46)
DIM    = RGBColor(0x8B, 0x8B, 0x93)
LINE   = RGBColor(0xE4, 0xE4, 0xE7)
LINE_S = RGBColor(0xD4, 0xD4, 0xD8)
CG     = RGBColor(235, 123, 22)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

# ---------- fonty ----------
# Arial: obecny na kazdej maszynie, wiec PowerPoint niczego nie podmieni.
# Role "monospace" z CSS (eyebrow, kicker, naglowki tabel, numer slajdu) zostaja
# rozpoznawalne przez wersaliki + rozstrzelenie + bold, tak jak w oryginale.
SANS = "Arial"
CODE = "Courier New"
ARIAL = {
    (False, False): "/System/Library/Fonts/Supplemental/Arial.ttf",
    (True,  False): "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    (False, True ): "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
    (True,  True ): "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf",
}
_fc = {}
def _font(px, bold=False, italic=False):
    k = (round(px), bold, italic)
    if k not in _fc:
        _fc[k] = ImageFont.truetype(ARIAL[(bold, italic)], max(1, round(px)))
    return _fc[k]

WARN = []

# ---------- pomiar i zawijanie ----------

def _tokens(spans):
    for sp in spans:
        for i, part in enumerate(sp.text.split('\n')):
            if i:
                yield None, sp                      # twarde zlamanie
            for w in re.findall(r'\s*\S+\s*', part):
                yield w, sp

def wrap(spans, px, max_w, bold=False):
    """Zawija spany na szerokosc max_w (px plotna). Zwraca linie, linia to lista (tekst, span)."""
    lines, cur, cw = [], [], 0.0
    for tok, sp in _tokens(spans):
        if tok is None:
            lines.append(cur); cur, cw = [], 0.0; continue
        f = _font(px, bold or sp.bold, sp.italic)
        tw = f.getlength(tok)
        if cur and cw + f.getlength(tok.rstrip()) > max_w:
            lines.append(cur); cur, cw = [], 0.0
        cur.append((tok, sp)); cw += tw
    if cur:
        lines.append(cur)
    return lines or [[]]

def text_h(spans, px, max_w, lh, bold=False):
    return len(wrap(spans, px, max_w, bold)) * px * lh

# ---------- prymitywy ----------

def _spc(run, px, em):
    """letter-spacing w 1/100 pt, jak a:rPr/@spc w OOXML."""
    if not em:
        return
    run.font._rPr.set('spc', str(int(round(px * FS * 0.6 * em * 100))))

def box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return tf

def para(tf, spans, px, color, lh=1.5, bold=False, align=PP_ALIGN.LEFT,
         upper=False, tracking=0.0, first=True):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = Pt(round(px * FS * 0.6 * lh, 2))
    p.space_before = Pt(0); p.space_after = Pt(0)
    if not spans:
        spans = [ds.Span('')]
    for sp in spans:
        # pojedyncze zlamanie w akapicie to <br> w CSS. W OOXML goly \n wewnatrz <a:t>
        # nie jest zlamaniem, PowerPoint go znormalizuje - trzeba wstawic <a:br/>.
        for j, piece in enumerate(sp.text.split('\n')):
            if j:
                p.add_line_break()
            if not piece:
                continue
            r = p.add_run()
            r.text = piece.upper() if upper else piece
            r.font.size = P(px)
            r.font.name = CODE if sp.code else SANS
            r.font.bold = bool(bold or sp.bold)
            r.font.italic = bool(sp.italic)
            # strong w CSS zmienia kolor na --text, reszta zostaje w kolorze bloku
            r.font.color.rgb = TEXT if (sp.bold and not bold) else color
            if sp.href:
                r.hyperlink.address = sp.href
                r.font.color.rgb = CG
            _spc(r, px, tracking)
    return p

def rect(slide, x, y, w, h, color, radius=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        E(x), E(y), E(w), E(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    if radius:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    return shp

def picture(slide, src, bx, by, bw, bh):
    """object-fit: contain - skaluje z zachowaniem proporcji i centruje w ramce."""
    p = HERE / src
    from PIL import Image
    with Image.open(p) as im:
        iw, ih = im.size
    k = min(bw / iw, bh / ih)
    w, h = iw * k, ih * k
    return slide.shapes.add_picture(str(p), E(bx + (bw - w) / 2), E(by + (bh - h) / 2), E(w), E(h))

def cell_border(cell, edge, color, pt_w):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = f'a:ln{edge}'
    for old in tcPr.findall(qn(tag)):
        tcPr.remove(old)
    ln = tcPr.makeelement(qn(tag), {'w': str(int(pt_w * 12700)), 'cap': 'flat',
                                    'cmpd': 'sng', 'algn': 'ctr'})
    fill = ln.makeelement(qn('a:solidFill'), {})
    clr = ln.makeelement(qn('a:srgbClr'), {'val': str(color)})
    fill.append(clr); ln.append(fill)
    # kolejnosc elementow w tcPr jest narzucona przez schemat: lnL, lnR, lnT, lnB
    order = ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB']
    idx = order.index(tag)
    anchor = None
    for later in order[idx + 1:]:
        found = tcPr.find(qn(later))
        if found is not None:
            anchor = found; break
    tcPr.insert(list(tcPr).index(anchor) if anchor is not None else 0, ln)

# ---------- bloki tresci ----------

def block_h(b, px, w, lh, gap_li=16):
    if b.kind == 'p':
        size = {'lede': 21, 'ask': 44}.get(b.cls, px)
        h = text_h(b.spans, SZ(size), w, {'lede': 1.5, 'ask': 1.15}.get(b.cls, lh),
                   bold=b.cls == 'ask')
        return h + (22 + 5 if b.cls == 'ask' else 0)
    if b.kind in ('ul', 'ol'):
        return sum(text_h(it, SZ(px), w - 34, 1.45) for it in b.items) + gap_li * (len(b.items) - 1)
    if b.kind == 'table':
        hh = SZ(12) * 1.3 + 11 + 2
        return hh + sum(SZ(19) * 1.4 + 24 for _ in b.rows)
    return 0

def draw_block(slide, b, x, y, w, px, lh):
    if b.kind == 'p':
        size, color, bold, llh = px, MUTED, False, lh
        if b.cls == 'lede': size, color = 21, DIM
        if b.cls == 'ask':  size, color, bold, llh = 44, TEXT, True, 1.15
        h = text_h(b.spans, SZ(size), w, llh, bold)
        tf = box(slide, x, y, w, h + 4)
        para(tf, b.spans, size, color, llh, bold=bold)
        if b.cls == 'ask':
            rect(slide, x, y + h + 22, 74, 5, CG, radius=3)
            return h + 22 + 5
        return h

    if b.kind in ('ul', 'ol'):
        cy = y
        for i, it in enumerate(b.items, 1):
            ih = text_h(it, SZ(px), w - 34, 1.45)
            tf = box(slide, x + 34, cy, w - 34, ih + 4)
            para(tf, it, px, MUTED, 1.45)
            if b.kind == 'ul':
                rect(slide, x + 2, cy + SZ(px) * 0.52, 12, 3, CG, radius=2)
            else:
                nt = box(slide, x, cy + SZ(px) * 0.06, 30, SZ(19) * 1.6)
                para(nt, [ds.Span(str(i))], 19, CG, 1.2, bold=True)
            cy += ih + 16
        return cy - y - 16

    if b.kind == 'table':
        ncol = len(b.head)
        nat = []
        for c in range(ncol):
            wmax = _font(SZ(12), True).getlength(ds.plain(b.head[c]).upper()) + SZ(12) * 0.13 * len(ds.plain(b.head[c]))
            for r in b.rows:
                if c < len(r):
                    wmax = max(wmax, _font(SZ(19), True).getlength(ds.plain(r[c])))
            nat.append(wmax + 20)
        k = w / sum(nat)
        widths = [n * k for n in nat]

        hh = SZ(12) * 1.3 + 11 + 2
        rh = SZ(19) * 1.4 + 24
        shp = slide.shapes.add_table(len(b.rows) + 1, ncol, E(x), E(y), E(w), E(hh + rh * len(b.rows)))
        tbl = shp.table
        tbl.first_row = False; tbl.horz_banding = False
        for i, cw in enumerate(widths):
            tbl.columns[i].width = E(cw)
        tbl.rows[0].height = E(hh)
        for i in range(len(b.rows)):
            tbl.rows[i + 1].height = E(rh)

        for c in range(ncol):
            cell = tbl.cell(0, c)
            cell.fill.background()
            cell.margin_left = cell.margin_top = Emu(0)
            cell.margin_right = E(20 if c < ncol - 1 else 0); cell.margin_bottom = E(11)
            cell.vertical_anchor = MSO_ANCHOR.BOTTOM
            cell.text_frame.word_wrap = True
            para(cell.text_frame, b.head[c], 12, DIM, 1.3, bold=True, upper=True, tracking=0.13)
            for e in ('L', 'R', 'T'):
                cell_border(cell, e, 'FFFFFF', 0)
            cell_border(cell, 'B', 'D4D4D8', 2 * FS * 0.6)

        for r, row in enumerate(b.rows, 1):
            last = r == len(b.rows)
            for c in range(ncol):
                cell = tbl.cell(r, c)
                cell.fill.background()
                cell.margin_left = Emu(0); cell.margin_top = E(12)
                cell.margin_right = E(20 if c < ncol - 1 else 0); cell.margin_bottom = E(12)
                cell.vertical_anchor = MSO_ANCHOR.TOP
                cell.text_frame.word_wrap = True
                para(cell.text_frame, row[c] if c < len(row) else [], 19, MUTED, 1.4)
                for e in ('L', 'R', 'T'):
                    cell_border(cell, e, 'FFFFFF', 0)
                cell_border(cell, 'B', 'FFFFFF' if last else 'E4E4E7', 0 if last else 1 * FS * 0.6)
        return hh + rh * len(b.rows)
    return 0

def flow(slide, blocks, x, y, w, h, px, lh, gap=24, label=''):
    """Uklada bloki w pionie, wysrodkowane w (y, h) - odpowiednik .body{justify-content:center}."""
    hs = [block_h(b, px, w, lh) for b in blocks]
    total = sum(hs) + gap * (len(blocks) - 1) if blocks else 0
    if total > h + 1:
        WARN.append(f"slajd {label}: tresc przekracza ramke o {round(total - h)} px")
    cy = y + max(0, (h - total) / 2)
    for b in blocks:
        cy += draw_block(slide, b, x, cy, w, px, lh) + gap

# ---------- naglowek ----------

def header(slide, s, h2_px, thin=False):
    """Zwraca y konca naglowka. Odwzorowuje .eyebrow + header h2 z CSS."""
    y = PAD_T
    if s.eyebrow:
        px = SZ(14)
        tw = _font(px, True).getlength(s.eyebrow.upper()) + px * 0.2 * len(s.eyebrow) + 4
        tf = box(slide, PAD_X, y, min(tw, CONTENT_W), px * 1.3)
        para(tf, [ds.Span(s.eyebrow)], 14, CG, 1.3, bold=True, upper=True, tracking=0.2)
        lx = PAD_X + min(tw, CONTENT_W) + 20
        if lx < W - PAD_X:
            rect(slide, lx, y + px * 0.62, W - PAD_X - lx, 1, LINE)
        y += px * 1.3 + 30
    if s.title and not thin:
        hh = text_h([ds.Span(s.title)], SZ(h2_px), min(1340, CONTENT_W), 1.08, bold=True)
        tf = box(slide, PAD_X, y, min(1340, CONTENT_W), hh + 6)
        para(tf, [ds.Span(s.title)], h2_px, TEXT, 1.08, bold=True)
        y += hh
    return y

def pnum(slide, n):
    px = SZ(14)
    tf = box(slide, W - PAD_X - 120, H - 34 - px * 1.2, 120, px * 1.4)
    para(tf, [ds.Span(str(n))], 14, DIM, 1.2, align=PP_ALIGN.RIGHT, tracking=0.08)

# ---------- slajdy ----------

def render(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])          # uklad pusty
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = WHITE
    blocks = ds.blocks(s.text)

    if s.kind == 'title':
        rect(slide, 0, 0, 14, H, CG)                            # .t-rule
        x, w = 110, W - 110 - PAD_X
        kh = SZ(17) * 1.3
        h1h = text_h([ds.Span(s.big or '')], SZ(118), w, 0.98, bold=True)
        bh = sum(block_h(b, 26, w, 1.5) for b in blocks) + 24 * max(0, len(blocks) - 1)
        total = kh + 28 + h1h + 38 + bh
        y = (H - total) / 2
        if s.title:
            para(box(slide, x, y, w, kh), [ds.Span(s.title)], 17, CG, 1.3,
                 bold=True, upper=True, tracking=0.24)
        y += kh + 28
        para(box(slide, x, y, w, h1h + 8), [ds.Span(s.big or '')], 118, TEXT, 0.98, bold=True)
        y += h1h + 38
        for b in blocks:
            y += draw_block(slide, b, x, y, w, 26, 1.5) + 24

    elif s.kind == 'closing':
        w = CONTENT_W
        kh = SZ(17) * 1.3
        h1h = text_h([ds.Span(s.eyebrow or '')], SZ(96), w, 1.0, bold=True)
        bh = sum(block_h(b, 27, w, 1.65) for b in blocks) + 24 * max(0, len(blocks) - 1)
        total = kh + 24 + h1h + 44 + bh
        y = (H - total) / 2
        if s.title:
            para(box(slide, PAD_X, y, w, kh), [ds.Span(s.title)], 17, CG, 1.3,
                 bold=True, upper=True, tracking=0.24, align=PP_ALIGN.CENTER)
        y += kh + 24
        para(box(slide, PAD_X, y, w, h1h + 8), [ds.Span(s.eyebrow or '')], 96, TEXT, 1.0,
             bold=True, align=PP_ALIGN.CENTER)
        y += h1h + 44
        for b in blocks:
            hh = block_h(b, 27, w, 1.65)
            tf = box(slide, PAD_X, y, w, hh + 4)
            para(tf, b.spans, 27, MUTED, 1.65, align=PP_ALIGN.CENTER)
            y += hh + 24

    elif s.kind == 'figure':
        alt, src, _ = s.imgs[0]
        own = (HERE / (pathlib.Path(src).stem + '.html')).exists()   # infografika ma wlasny tytul
        y = header(slide, s, 31, thin=own) + (16 if own else 22)
        picture(slide, src, PAD_X, y, CONTENT_W, H - PAD_B - y)
        pnum(slide, s.index)

    elif s.kind == 'split':
        y = header(slide, s, 41) + 34
        h = H - PAD_B - y
        lw = (CONTENT_W - 52) / 1.95                              # grid 1fr 0.95fr, gap 52
        rw = CONTENT_W - 52 - lw
        flow(slide, blocks, PAD_X, y, lw, h, 25, 1.5, gap=20, label=str(s.index))
        _, src, _ = s.imgs[0]
        picture(slide, src, PAD_X + lw + 52, y, rw, h)
        pnum(slide, s.index)

    else:
        h2 = 64 if s.sparse else 41
        y = header(slide, s, h2) + (46 if s.sparse else 34)
        flow(slide, blocks, PAD_X, y, CONTENT_W, H - PAD_B - y,
             31 if s.sparse else 27, 1.45 if s.sparse else 1.5, label=str(s.index))
        pnum(slide, s.index)

    if s.notes:
        slide.notes_slide.notes_text_frame.text = s.notes
    return slide

# ---------- main ----------

def main():
    slides = ds.load()
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)
    for s in slides:
        render(prs, s)
    prs.save(OUT)

    notes = sum(1 for s in slides if s.notes)
    print(f"{len(slides)} slajdów -> {OUT.name}  ({notes} notatek prowadzącego)")
    if WARN:
        print("\n  UWAGA - tresc nie miesci sie w ramce:")
        for w in WARN:
            print("   -", w)
        print("  Skroc tekst w deck.md. Rozmiary sa celowo sztywne,")
        print("  zeby PowerPoint nie zmniejszal tekstu po swojemu.")
    return 0

if __name__ == "__main__":
    sys.exit(main())
