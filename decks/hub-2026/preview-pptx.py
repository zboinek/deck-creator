#!/usr/bin/env python3
"""Rasteryzuje .pptx do PNG czytajac zapisany plik - bez Keynote, LibreOffice i GUI.

Wszystko (pozycje, rozmiary fontow, kolory, obramowania tabel) pochodzi z XML-a
w .pptx, wiec to jest kontrola tego, co naprawde wyladowalo w pliku.
Uwaga: to moj renderer, nie PowerPointa - dowodzi poprawnosci zawartosci
i braku wyjscia poza ramke, nie identycznosci z PowerPointem co do piksela."""
import sys, pathlib
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

PXE = 7620.0
def px(emu): return emu / PXE
def fpx(pt): return pt / 0.6          # 0.6 pt na piksel plotna

ARIAL = {(0,0):"/System/Library/Fonts/Supplemental/Arial.ttf",
         (1,0):"/System/Library/Fonts/Supplemental/Arial Bold.ttf",
         (0,1):"/System/Library/Fonts/Supplemental/Arial Italic.ttf",
         (1,1):"/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf"}
_fc={}
def font(size, b=False, i=False):
    k=(round(size),bool(b),bool(i))
    if k not in _fc: _fc[k]=ImageFont.truetype(ARIAL[(int(bool(b)),int(bool(i)))], max(1,round(size)))
    return _fc[k]

def rgb(c, d=(0,0,0)):
    try: return (c.rgb[0], c.rgb[1], c.rgb[2])
    except Exception: return d

def runs_of(p):
    """Chodzi po dzieciach <a:p> po kolei, zeby zlapac tez <a:br/> - p.runs je pomija."""
    by_el = {r._r: r for r in p.runs}
    out = []
    for el in p._p:
        if el.tag == qn('a:br'):
            out.append(dict(br=True, t='', sz=0, b=False, i=False, c=(0,0,0), spc=0))
        elif el.tag == qn('a:r') and el in by_el:
            r = by_el[el]
            spc = 0
            v = r.font._rPr.get('spc')
            if v:
                spc = fpx(int(v) / 100.0)
            out.append(dict(br=False, t=r.text,
                            sz=fpx(r.font.size.pt) if r.font.size else 20,
                            b=bool(r.font.bold), i=bool(r.font.italic),
                            c=rgb(r.font.color), spc=spc))
    return out

def draw_text(d, tf, x, y, w, h, oob):
    for p in tf.paragraphs:
        rs = runs_of(p)
        if not rs: continue
        body = [r for r in rs if not r.get('br')]
        if not body: continue
        sz = max(r['sz'] for r in body)
        lh = fpx(p.line_spacing.pt) if hasattr(p.line_spacing,'pt') and p.line_spacing else sz*1.2
        # zawijanie na zapisanej szerokosci pola
        lines, cur, cw = [], [], 0.0
        for r in rs:
            if r.get('br'):
                lines.append(cur); cur,cw=[],0.0; continue
            for tok in __import__('re').findall(r'\n|[^\S\n]*\S+[^\S\n]*', r['t']):
                if tok=='\n': lines.append(cur); cur,cw=[],0.0; continue
                f=font(r['sz'],r['b'],r['i'])
                tw=f.getlength(tok)+r['spc']*len(tok)
                if cur and cw+f.getlength(tok.rstrip())+r['spc']*len(tok.rstrip())>w:
                    lines.append(cur); cur,cw=[],0.0
                cur.append((tok,r)); cw+=tw
        if cur: lines.append(cur)
        for ln in lines:
            ln=[(t_,r_) for t_,r_ in ln if not r_.get('br')]
            lw=sum(font(r['sz'],r['b'],r['i']).getlength(t_)+r['spc']*len(t_) for t_,r in ln)
            cx = x + (w-lw)/2 if p.alignment==PP_ALIGN.CENTER else \
                 x + (w-lw)   if p.alignment==PP_ALIGN.RIGHT  else x
            for t_,r in ln:
                t=t_
                f=font(r['sz'],r['b'],r['i'])
                if r['spc']:
                    for ch in t:
                        d.text((cx,y),ch,font=f,fill=r['c']); cx+=f.getlength(ch)+r['spc']
                else:
                    d.text((cx,y),t,font=f,fill=r['c']); cx+=f.getlength(t)
            y+=lh
            if y > 900: oob.append('tekst ponizej dolu slajdu')
        if y > 900: oob.append('akapit wychodzi poza slajd')
    return y

def border_b(cell):
    tcPr=cell._tc.find(qn('a:tcPr'))
    if tcPr is None: return None
    ln=tcPr.find(qn('a:lnB'))
    if ln is None: return None
    clr=ln.find(qn('a:solidFill/a:srgbClr')) if False else ln.find(qn('a:solidFill'))
    val=clr.find(qn('a:srgbClr')).get('val') if clr is not None else '000000'
    return int(ln.get('w',0))/12700.0, tuple(int(val[i:i+2],16) for i in (0,2,4))

def render(prs, idx, slide, outdir):
    img=Image.new('RGB',(1600,900),'white'); d=ImageDraw.Draw(img); oob=[]
    for shp in slide.shapes:
        x,y,w,h = px(shp.left), px(shp.top), px(shp.width), px(shp.height)
        if x<-1 or y<-1 or x+w>1601 or y+h>901:
            oob.append(f'ksztalt poza slajdem: {round(x)},{round(y)} {round(w)}x{round(h)}')
        if shp.shape_type==13:                                   # obrazek
            im=Image.open(__import__('io').BytesIO(shp.image.blob)).convert('RGBA')
            im=im.resize((max(1,round(w)),max(1,round(h))))
            img.paste(im,(round(x),round(y)),im)
        elif shp.has_table:
            t=shp.table; cy=y
            for ri,row in enumerate(t.rows):
                cx=x
                for ci,cell in enumerate(row.cells):
                    cw=px(t.columns[ci].width); rh=px(row.height)
                    ml,mt,mr=px(cell.margin_left),px(cell.margin_top),px(cell.margin_right)
                    ty=cy+mt
                    if str(cell.vertical_anchor)=='BOTTOM (5)' or cell.vertical_anchor is not None and 'BOTTOM' in str(cell.vertical_anchor):
                        ty=cy+rh-px(cell.margin_bottom)-fpx(12*0.85*0.6)*1.3
                    draw_text(d,cell.text_frame,cx+ml,ty,cw-ml-mr,rh,oob)
                    b=border_b(cell)
                    if b and b[0]>0:
                        lw=max(1,round(fpx(b[0])))
                        d.rectangle([cx,cy+rh-lw,cx+cw,cy+rh],fill=b[1])
                    cx+=cw
                cy+=px(row.height)
        elif shp.has_text_frame and shp.text_frame.text.strip():
            if shp.shape_type==1 and shp.fill.type is not None:
                pass
            draw_text(d,shp.text_frame,x,y,w,h,oob)
        elif shp.shape_type==1:                                  # autoksztalt bez tekstu
            try: col=rgb(shp.fill.fore_color,(200,200,200))
            except Exception: col=(200,200,200)
            d.rounded_rectangle([x,y,x+w,y+h],radius=min(h/2,3),fill=col)
    img.save(outdir/f'slide-{idx:02d}.png')
    return oob

def main():
    src=pathlib.Path(sys.argv[1]); outdir=pathlib.Path(sys.argv[2]); outdir.mkdir(parents=True,exist_ok=True)
    prs=Presentation(str(src))
    problems={}
    for i,s in enumerate(prs.slides,1):
        o=render(prs,i,s,outdir)
        if o: problems[i]=sorted(set(o))
    print(f"wyrenderowano {len(prs.slides)} slajdow -> {outdir}")
    if problems:
        print("\nPROBLEMY GEOMETRII:")
        for k,v in problems.items():
            for p in v: print(f"  slajd {k}: {p}")
    else:
        print("geometria czysta: zaden ksztalt ani wiersz tekstu nie wychodzi poza slajd")

main()
